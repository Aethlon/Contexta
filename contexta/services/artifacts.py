from __future__ import annotations

import hashlib
import inspect
import io
import json
import logging
import mimetypes
import os
import tempfile
from collections.abc import Mapping
from dataclasses import dataclass
from email.message import Message
from html.parser import HTMLParser
from pathlib import Path, PurePosixPath
from typing import Any, Protocol
from urllib.parse import unquote, urlparse
from uuid import UUID, uuid4

import httpx

from contexta.repositories.artifact_repo import ArtifactRepository

DEFAULT_MAX_ARTIFACT_SIZE_BYTES = 25 * 1024 * 1024
DEFAULT_LOCAL_STORAGE_PATH = "data/artifacts"
DEFAULT_FETCH_TIMEOUT_SECONDS = 60.0
EXTRACTION_COMPLETED_STATUS = "completed"
SUCCESSFUL_EXTRACTION_STATUSES = frozenset({"completed", "ready", "extracted"})

logger = logging.getLogger(__name__)


class ArtifactError(Exception):
    pass


class ArtifactCapabilityError(ArtifactError):
    pass


class ArtifactValidationError(ArtifactError):
    pass


class ArtifactSizeError(ArtifactValidationError):
    pass


class ArtifactStorageError(ArtifactError):
    def __init__(self, message: str, *, operation: str | None = None) -> None:
        self.operation = operation
        super().__init__(message)


class ArtifactDownloadError(ArtifactError):
    pass


class ArtifactExtractionError(ArtifactError):
    pass


class ArtifactNotFoundError(ArtifactError):
    pass


@dataclass(frozen=True, slots=True)
class StoredObject:
    backend: str
    key: str
    reference: str
    etag: str | None = None


@dataclass(frozen=True, slots=True)
class ExtractionResult:
    text: str
    method: str | None
    status: str
    error: str | None = None


@dataclass(frozen=True, slots=True)
class FetchedURL:
    data: bytes
    filename: str
    mime_type: str
    original_url: str
    resolved_url: str


class ArtifactStorage(Protocol):
    backend_name: str

    async def put(
        self,
        *,
        key: str,
        data: bytes,
        content_type: str,
        metadata: Mapping[str, str],
    ) -> StoredObject: ...

    async def get(self, *, key: str) -> bytes: ...

    async def delete(self, *, key: str) -> None: ...


def _normalize_storage_key(key: str) -> str:
    value = str(key).replace("\\", "/").strip("/")
    if not value or "\x00" in value or ":" in value:
        raise ArtifactValidationError("Invalid artifact storage key.")
    parts = value.split("/")
    if any(not part or part in {".", ".."} for part in parts):
        raise ArtifactValidationError("Invalid artifact storage key.")
    return "/".join(parts)


class LocalArtifactStorage:
    backend_name = "local"

    def __init__(self, root: str | Path | None = None) -> None:
        configured_root = root or os.environ.get(
            "CONTEXTA_ARTIFACT_STORAGE_PATH", DEFAULT_LOCAL_STORAGE_PATH
        )
        self._root = Path(configured_root).expanduser().resolve()

    def _path(self, key: str) -> Path:
        normalized = _normalize_storage_key(key)
        candidate = (self._root / Path(*PurePosixPath(normalized).parts)).resolve()
        if candidate == self._root or self._root not in candidate.parents:
            raise ArtifactValidationError("Invalid artifact storage key.")
        return candidate

    async def put(
        self,
        *,
        key: str,
        data: bytes,
        content_type: str,
        metadata: Mapping[str, str],
    ) -> StoredObject:
        del content_type, metadata
        target = self._path(key)
        temporary_path: Path | None = None
        try:
            target.parent.mkdir(parents=True, exist_ok=True)
            with tempfile.NamedTemporaryFile(
                mode="wb",
                dir=target.parent,
                prefix=".artifact-",
                delete=False,
            ) as handle:
                temporary_path = Path(handle.name)
                handle.write(data)
                handle.flush()
                os.fsync(handle.fileno())
            os.replace(temporary_path, target)
        except OSError as exc:
            raise ArtifactStorageError(
                "Local artifact storage could not write the original.",
                operation="put",
            ) from exc
        finally:
            if temporary_path is not None:
                temporary_path.unlink(missing_ok=True)
        return StoredObject(
            backend=self.backend_name,
            key=key,
            reference=f"local://{key}",
        )

    async def get(self, *, key: str) -> bytes:
        target = self._path(key)
        try:
            return target.read_bytes()
        except OSError as exc:
            raise ArtifactStorageError(
                "Local artifact storage could not read the original.",
                operation="get",
            ) from exc

    async def delete(self, *, key: str) -> None:
        target = self._path(key)
        try:
            target.unlink(missing_ok=True)
        except OSError as exc:
            raise ArtifactStorageError(
                "Local artifact storage could not delete the original.",
                operation="delete",
            ) from exc


class S3CompatibleStorage:
    backend_name = "s3"

    def __init__(self, client: Any, bucket: str, prefix: str = "") -> None:
        if not bucket.strip():
            raise ArtifactValidationError("An S3-compatible bucket is required.")
        for method_name in ("put_object", "get_object", "delete_object"):
            if not callable(getattr(client, method_name, None)):
                raise ArtifactValidationError(
                    "The S3-compatible client does not implement the required operations."
                )
        self._client = client
        self._bucket = bucket
        self._prefix = prefix.strip("/")

    def _key(self, key: str) -> str:
        normalized = _normalize_storage_key(key)
        if self._prefix:
            return f"{self._prefix}/{normalized}"
        return normalized

    async def _call(self, method_name: str, **kwargs: Any) -> Any:
        try:
            result = getattr(self._client, method_name)(**kwargs)
            if inspect.isawaitable(result):
                return await result
            return result
        except Exception as exc:
            operation = method_name.removesuffix("_object")
            raise ArtifactStorageError(
                f"S3-compatible storage failed during {operation}.",
                operation=operation,
            ) from exc

    async def put(
        self,
        *,
        key: str,
        data: bytes,
        content_type: str,
        metadata: Mapping[str, str],
    ) -> StoredObject:
        object_key = self._key(key)
        result = await self._call(
            "put_object",
            Bucket=self._bucket,
            Key=object_key,
            Body=data,
            ContentType=content_type,
            Metadata=dict(metadata),
        )
        etag = self._response_value(result, "ETag")
        if etag is not None:
            etag = str(etag)
        return StoredObject(
            backend=self.backend_name,
            key=key,
            reference=f"s3://{self._bucket}/{object_key}",
            etag=etag,
        )

    async def get(self, *, key: str) -> bytes:
        object_key = self._key(key)
        result = await self._call(
            "get_object", Bucket=self._bucket, Key=object_key
        )
        body = self._response_value(result, "Body")
        if body is None:
            raise ArtifactStorageError(
                "S3-compatible storage returned an empty object.", operation="get"
            )
        if isinstance(body, bytes):
            return body
        if isinstance(body, bytearray):
            return bytes(body)
        reader = getattr(body, "read", None)
        if not callable(reader):
            raise ArtifactStorageError(
                "S3-compatible storage returned an unreadable object.", operation="get"
            )
        try:
            value = reader()
            if inspect.isawaitable(value):
                value = await value
            return bytes(value)
        except Exception as exc:
            raise ArtifactStorageError(
                "S3-compatible storage could not read the object.", operation="get"
            ) from exc

    async def delete(self, *, key: str) -> None:
        await self._call(
            "delete_object", Bucket=self._bucket, Key=self._key(key)
        )

    @staticmethod
    def _response_value(result: Any, name: str) -> Any:
        if isinstance(result, Mapping):
            return result.get(name)
        return getattr(result, name, None)


LocalFileStorage = LocalArtifactStorage
S3ArtifactStorage = S3CompatibleStorage


def create_artifact_storage() -> ArtifactStorage:
    backend = os.environ.get(
        "CONTEXTA_ARTIFACT_STORAGE_BACKEND",
        os.environ.get("CONTEXTA_ARTIFACT_STORAGE", "local"),
    ).strip().lower()
    if backend in {"local", "filesystem"}:
        return LocalArtifactStorage()
    if backend in {"s3", "s3-compatible", "minio"}:
        try:
            import boto3
        except ImportError as exc:
            raise ArtifactCapabilityError(
                "S3-compatible artifact storage requires an injected client or the optional 'boto3' package."
            ) from exc
        bucket = os.environ.get("CONTEXTA_S3_BUCKET", "").strip()
        if not bucket:
            raise ArtifactValidationError("CONTEXTA_S3_BUCKET is required for S3 storage.")
        access_key = os.environ.get("CONTEXTA_S3_ACCESS_KEY_ID", "")
        secret_key = os.environ.get("CONTEXTA_S3_SECRET_ACCESS_KEY", "")
        if bool(access_key) != bool(secret_key):
            raise ArtifactValidationError(
                "Both CONTEXTA_S3_ACCESS_KEY_ID and CONTEXTA_S3_SECRET_ACCESS_KEY are required."
            )
        client_kwargs: dict[str, Any] = {
            "service_name": "s3",
            "endpoint_url": os.environ.get("CONTEXTA_S3_ENDPOINT_URL") or None,
            "region_name": os.environ.get("CONTEXTA_S3_REGION") or None,
        }
        if access_key and secret_key:
            client_kwargs["aws_access_key_id"] = access_key
            client_kwargs["aws_secret_access_key"] = secret_key
        try:
            client = boto3.client(**client_kwargs)
        except Exception as exc:
            raise ArtifactStorageError(
                "S3-compatible artifact storage could not be initialized.", operation="init"
            ) from exc
        return S3CompatibleStorage(
            client=client,
            bucket=bucket,
            prefix=os.environ.get("CONTEXTA_S3_PREFIX", "artifacts"),
        )
    raise ArtifactValidationError(f"Unsupported artifact storage backend: {backend}.")


class _HTMLTextParser(HTMLParser):
    _SKIPPED_TAGS = frozenset({"script", "style", "noscript", "template"})

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self._parts: list[str] = []
        self._skip_depth = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        del attrs
        lowered = tag.casefold()
        if lowered in self._SKIPPED_TAGS:
            self._skip_depth += 1
        elif not self._skip_depth and lowered in {
            "br",
            "div",
            "li",
            "p",
            "section",
            "tr",
        }:
            self._parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        lowered = tag.casefold()
        if lowered in self._SKIPPED_TAGS and self._skip_depth:
            self._skip_depth -= 1
        elif not self._skip_depth and lowered in {
            "div",
            "li",
            "p",
            "section",
            "tr",
        }:
            self._parts.append("\n")

    def handle_data(self, data: str) -> None:
        if not self._skip_depth:
            self._parts.append(data)

    def text(self) -> str:
        return _normalize_text("".join(self._parts))


def _normalize_text(value: str) -> str:
    lines = [line.rstrip() for line in value.replace("\r\n", "\n").replace("\r", "\n").split("\n")]
    return "\n".join(lines).strip()


def _base_mime_type(value: str | None) -> str:
    if not value:
        return "application/octet-stream"
    return value.split(";", 1)[0].strip().lower() or "application/octet-stream"


def _canonical_extraction_status(status: str) -> str:
    value = str(status or "").strip().casefold()
    if value in SUCCESSFUL_EXTRACTION_STATUSES:
        return EXTRACTION_COMPLETED_STATUS
    return value


def _extension(filename: str) -> str:
    return Path(filename).suffix.casefold()


class ArtifactTextExtractor:
    _PLAIN_MIME_TYPES = frozenset(
        {
            "application/javascript",
            "application/json",
            "application/sql",
            "application/toml",
            "application/x-ndjson",
            "application/x-sh",
            "application/x-yaml",
            "application/xml",
            "application/yaml",
        }
    )
    _PLAIN_EXTENSIONS = frozenset(
        {
            ".c",
            ".cc",
            ".cpp",
            ".csv",
            ".css",
            ".go",
            ".h",
            ".hpp",
            ".ini",
            ".java",
            ".js",
            ".json",
            ".jsonl",
            ".jsx",
            ".log",
            ".lua",
            ".md",
            ".markdown",
            ".py",
            ".rb",
            ".rs",
            ".sh",
            ".sql",
            ".text",
            ".toml",
            ".ts",
            ".tsx",
            ".txt",
            ".xml",
            ".yaml",
            ".yml",
        }
    )

    def __init__(self, http_client: httpx.AsyncClient | None = None) -> None:
        self._http_client = http_client

    async def extract(
        self,
        *,
        data: bytes,
        filename: str,
        mime_type: str | None,
        source_url: str | None = None,
    ) -> ExtractionResult:
        del source_url
        mime = _base_mime_type(mime_type)
        extension = _extension(filename)
        if mime == "application/pdf" or extension == ".pdf":
            return self._extract_pdf(data)
        if mime == "application/vnd.openxmlformats-officedocument.wordprocessingml.document" or extension == ".docx":
            return self._extract_docx(data)
        if mime in {"text/html", "application/xhtml+xml"} or extension in {".html", ".htm", ".xhtml"}:
            parser = _HTMLTextParser()
            try:
                parser.feed(self._decode_text(data, mime_type))
                parser.close()
            except Exception as exc:
                raise ArtifactExtractionError("HTML text extraction failed.") from exc
            return ExtractionResult(parser.text(), "html", EXTRACTION_COMPLETED_STATUS)
        if (
            mime.startswith("text/")
            or mime in self._PLAIN_MIME_TYPES
            or extension in self._PLAIN_EXTENSIONS
            or not mime
        ):
            text = self._decode_text(data, mime_type)
            return ExtractionResult(
                _normalize_text(text),
                "plain_text",
                EXTRACTION_COMPLETED_STATUS,
            )
        return ExtractionResult("", None, "unsupported")

    async def fetch_url(
        self,
        *,
        url: str,
        max_size_bytes: int,
    ) -> FetchedURL:
        parsed = urlparse(url)
        if parsed.scheme.casefold() not in {"http", "https"} or not parsed.hostname:
            raise ArtifactValidationError("Artifact URLs must use HTTP or HTTPS.")
        if parsed.username or parsed.password:
            raise ArtifactValidationError("Artifact URLs cannot contain credentials.")
        if self._http_client is not None:
            return await self._download(
                self._http_client,
                url=url,
                max_size_bytes=max_size_bytes,
            )
        timeout = httpx.Timeout(DEFAULT_FETCH_TIMEOUT_SECONDS)
        async with httpx.AsyncClient(
            follow_redirects=True,
            timeout=timeout,
            headers={"User-Agent": "Contexta-Artifact/1.0"},
        ) as client:
            return await self._download(
                client, url=url, max_size_bytes=max_size_bytes
            )

    async def _download(
        self,
        client: httpx.AsyncClient,
        *,
        url: str,
        max_size_bytes: int,
    ) -> FetchedURL:
        try:
            async with client.stream("GET", url, follow_redirects=True) as response:
                if response.status_code < 200 or response.status_code >= 300:
                    raise ArtifactDownloadError(
                        f"Artifact URL returned HTTP {response.status_code}."
                    )
                content_length = response.headers.get("content-length")
                if content_length:
                    try:
                        if int(content_length) > max_size_bytes:
                            raise ArtifactSizeError(
                                f"Artifact exceeds the maximum size of {max_size_bytes} bytes."
                            )
                    except ValueError:
                        pass
                chunks: list[bytes] = []
                size = 0
                async for chunk in response.aiter_bytes():
                    size += len(chunk)
                    if size > max_size_bytes:
                        raise ArtifactSizeError(
                            f"Artifact exceeds the maximum size of {max_size_bytes} bytes."
                        )
                    chunks.append(chunk)
                disposition = response.headers.get("content-disposition", "")
                filename = self._filename_from_disposition(disposition)
                if not filename:
                    filename = self._filename_from_url(str(response.url))
                mime_type = response.headers.get(
                    "content-type", "application/octet-stream"
                )
                return FetchedURL(
                    data=b"".join(chunks),
                    filename=filename or "artifact",
                    mime_type=mime_type,
                    original_url=url,
                    resolved_url=str(response.url),
                )
        except (ArtifactError, httpx.HTTPError):
            raise
        except Exception as exc:
            raise ArtifactDownloadError("Artifact URL could not be downloaded.") from exc

    @staticmethod
    def _decode_text(data: bytes, mime_type: str | None) -> str:
        parameters: dict[str, str] = {}
        if mime_type and ";" in mime_type:
            for part in mime_type.split(";")[1:]:
                if "=" in part:
                    name, value = part.split("=", 1)
                    parameters[name.strip().casefold()] = value.strip().strip('"')
        encoding = parameters.get("charset")
        if not encoding and data.startswith((b"\xff\xfe", b"\xfe\xff")):
            encoding = "utf-16"
        try:
            return data.decode(encoding or "utf-8-sig")
        except (LookupError, UnicodeDecodeError) as exc:
            raise ArtifactExtractionError(
                "Plain text is not valid in the declared character encoding."
            ) from exc

    @staticmethod
    def _filename_from_disposition(value: str) -> str | None:
        if not value:
            return None
        message = Message()
        message["content-disposition"] = value
        return message.get_filename()

    @staticmethod
    def _filename_from_url(value: str) -> str | None:
        path = urlparse(value).path
        name = unquote(Path(path).name)
        return name or None

    @staticmethod
    def _extract_pdf(data: bytes) -> ExtractionResult:
        try:
            try:
                from pypdf import PdfReader
            except ImportError:
                from PyPDF2 import PdfReader
        except ImportError as exc:
            raise ArtifactCapabilityError(
                "PDF text extraction requires the 'pypdf' or 'PyPDF2' package."
            ) from exc
        try:
            reader = PdfReader(io.BytesIO(data))
            text = "\n\n".join(page.extract_text() or "" for page in reader.pages)
        except Exception as exc:
            raise ArtifactExtractionError("PDF text extraction failed.") from exc
        return ExtractionResult(
            _normalize_text(text),
            "pdf",
            EXTRACTION_COMPLETED_STATUS,
        )

    @staticmethod
    def _extract_docx(data: bytes) -> ExtractionResult:
        try:
            import docx
        except ImportError as exc:
            raise ArtifactCapabilityError(
                "DOCX text extraction requires the 'python-docx' package."
            ) from exc
        try:
            document = docx.Document(io.BytesIO(data))
            parts = [paragraph.text for paragraph in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    parts.append("\t".join(cell.text for cell in row.cells))
        except Exception as exc:
            raise ArtifactExtractionError("DOCX text extraction failed.") from exc
        return ExtractionResult(
            _normalize_text("\n".join(parts)),
            "docx",
            EXTRACTION_COMPLETED_STATUS,
        )


class ArtifactService:
    def __init__(
        self,
        repository: ArtifactRepository,
        *,
        storage: ArtifactStorage | None = None,
        extractor: ArtifactTextExtractor | None = None,
        http_client: httpx.AsyncClient | None = None,
        max_size_bytes: int | None = None,
    ) -> None:
        self.repository = repository
        self._storage = storage or create_artifact_storage()
        self._extractor = extractor or ArtifactTextExtractor(http_client=http_client)
        self._max_size_bytes = self._read_max_size(max_size_bytes)

    @property
    def max_size_bytes(self) -> int:
        return self._max_size_bytes

    async def store_artifact(
        self,
        *,
        data: bytes,
        filename: str,
        user_id: UUID,
        mime_type: str | None = None,
        original_reference: str | None = None,
        metadata: Mapping[str, Any] | None = None,
        artifact_id: UUID | None = None,
    ) -> Any:
        original = bytes(data)
        if len(original) > self._max_size_bytes:
            raise ArtifactSizeError(
                f"Artifact exceeds the maximum size of {self._max_size_bytes} bytes."
            )
        safe_filename = self._normalize_filename(filename)
        resolved_mime_type = self._resolve_mime_type(safe_filename, mime_type)
        extraction_mime_type = mime_type or resolved_mime_type
        if _base_mime_type(extraction_mime_type) == "application/octet-stream":
            extraction_mime_type = resolved_mime_type
        metadata_value = self._normalize_metadata(metadata)
        identifier = artifact_id or uuid4()
        storage_key = self._storage_key(identifier, safe_filename)

        try:
            extraction = await self._extractor.extract(
                data=original,
                filename=safe_filename,
                mime_type=extraction_mime_type,
                source_url=original_reference,
            )
        except ArtifactCapabilityError as exc:
            extraction = ExtractionResult("", None, "capability_unavailable", str(exc))
        except ArtifactExtractionError as exc:
            extraction = ExtractionResult("", None, "failed", str(exc))
        except (OSError, ValueError, TypeError, UnicodeError, ImportError, RuntimeError) as exc:
            extraction = ExtractionResult(
                "",
                None,
                "failed",
                f"{type(exc).__name__}: {str(exc)[:900]}",
            )

        content_hash = hashlib.sha256(original).hexdigest()
        stored = await self._storage.put(
            key=storage_key,
            data=original,
            content_type=resolved_mime_type,
            metadata={
                "sha256": content_hash,
                "filename": safe_filename,
            },
        )
        try:
            return await self.repository.create_artifact(
                user_id=user_id,
                filename=safe_filename,
                mime_type=resolved_mime_type,
                size_bytes=len(original),
                content_hash=content_hash,
                storage_backend=stored.backend,
                storage_key=stored.key,
                storage_etag=stored.etag,
                original_reference=(original_reference or stored.reference).strip(),
                metadata=metadata_value,
                extracted_text=extraction.text,
                extraction_method=extraction.method,
                extraction_status=_canonical_extraction_status(extraction.status),
                extraction_error=extraction.error,
            )
        except BaseException:
            logger.exception(
                "Artifact metadata persistence failed; original retained",
                extra={
                    "artifact_id": str(identifier),
                    "storage_backend": stored.backend,
                    "storage_key": stored.key,
                    "content_hash": content_hash,
                },
            )
            raise

    async def create_artifact(self, **kwargs: Any) -> Any:
        return await self.store_artifact(**kwargs)

    async def create_from_url(
        self,
        *,
        url: str,
        user_id: UUID,
        filename: str | None = None,
        mime_type: str | None = None,
        original_reference: str | None = None,
        metadata: Mapping[str, Any] | None = None,
    ) -> Any:
        fetched = await self._extractor.fetch_url(
            url=url, max_size_bytes=self._max_size_bytes
        )
        metadata_value = dict(metadata or {})
        metadata_value.setdefault("resolved_url", fetched.resolved_url)
        return await self.store_artifact(
            data=fetched.data,
            filename=filename or fetched.filename,
            user_id=user_id,
            mime_type=mime_type or fetched.mime_type,
            original_reference=original_reference or fetched.original_url,
            metadata=metadata_value,
        )

    async def list_metadata(
        self,
        *,
        user_id: UUID | None = None,
        offset: int = 0,
        limit: int = 100,
    ) -> list[Any]:
        return list(
            await self.repository.list_artifacts(
                user_id=user_id, offset=offset, limit=limit
            )
        )

    async def get_metadata(self, artifact_id: UUID) -> Any:
        record = await self.repository.get_by_id(artifact_id)
        if record is None:
            raise ArtifactNotFoundError("Artifact not found.")
        return record

    async def read_original(self, artifact_id: UUID) -> tuple[Any, bytes]:
        record = await self.get_metadata(artifact_id)
        return record, await self._storage.get(key=record.storage_key)

    async def reextract_text(self, artifact_id: UUID) -> Any:
        record, data = await self.read_original(artifact_id)
        extraction = await self._extractor.extract(
            data=data,
            filename=record.filename,
            mime_type=record.mime_type,
            source_url=record.original_reference,
        )
        values = {
            "extracted_text": extraction.text,
            "extraction_method": extraction.method,
            "extraction_status": _canonical_extraction_status(extraction.status),
            "extraction_error": extraction.error,
        }
        await self.repository.update_by_id(artifact_id, values)
        for key, value in values.items():
            setattr(record, key, value)
        return record

    @staticmethod
    def _normalize_filename(filename: str) -> str:
        value = Path(str(filename).replace("\\", "/")).name
        value = "".join(character for character in value if ord(character) >= 32)
        value = value.strip().strip(".")
        if not value:
            value = "artifact"
        return value[:500]

    @staticmethod
    def _resolve_mime_type(filename: str, supplied: str | None) -> str:
        resolved = _base_mime_type(supplied)
        guessed = mimetypes.guess_type(filename)[0]
        if guessed == "text/markdown":
            return "text/markdown"
        if resolved == "application/octet-stream" and guessed:
            return _base_mime_type(guessed)
        return resolved

    @staticmethod
    def _normalize_metadata(
        metadata: Mapping[str, Any] | None,
    ) -> dict[str, Any]:
        if metadata is None:
            return {}
        if not isinstance(metadata, Mapping):
            raise ArtifactValidationError("Artifact metadata must be an object.")
        value = dict(metadata)
        try:
            encoded = json.dumps(value, separators=(",", ":"), default=str).encode(
                "utf-8"
            )
        except (TypeError, ValueError) as exc:
            raise ArtifactValidationError("Artifact metadata must be JSON serializable.") from exc
        if len(encoded) > 65_536:
            raise ArtifactValidationError("Artifact metadata exceeds 64 KiB.")
        return value

    @staticmethod
    def _storage_key(identifier: UUID, filename: str) -> str:
        extension = _extension(filename)
        if len(extension) > 20 or any(
            character not in "abcdefghijklmnopqrstuvwxyz0123456789." for character in extension
        ):
            extension = ""
        return f"{identifier}/{identifier}{extension}"

    @staticmethod
    def _read_max_size(value: int | None) -> int:
        if value is not None:
            if value <= 0:
                raise ArtifactValidationError("Maximum artifact size must be positive.")
            return value
        raw = os.environ.get("CONTEXTA_MAX_ARTIFACT_SIZE_BYTES", "")
        if raw:
            try:
                parsed = int(raw)
            except ValueError:
                parsed = 0
            if parsed > 0:
                return parsed
        return DEFAULT_MAX_ARTIFACT_SIZE_BYTES


__all__ = [
    "DEFAULT_MAX_ARTIFACT_SIZE_BYTES",
    "EXTRACTION_COMPLETED_STATUS",
    "SUCCESSFUL_EXTRACTION_STATUSES",
    "ArtifactCapabilityError",
    "ArtifactDownloadError",
    "ArtifactError",
    "ArtifactExtractionError",
    "ArtifactNotFoundError",
    "ArtifactService",
    "ArtifactSizeError",
    "ArtifactStorage",
    "ArtifactStorageError",
    "ArtifactTextExtractor",
    "ArtifactValidationError",
    "ExtractionResult",
    "FetchedURL",
    "LocalArtifactStorage",
    "LocalFileStorage",
    "S3ArtifactStorage",
    "S3CompatibleStorage",
    "StoredObject",
    "create_artifact_storage",
]
